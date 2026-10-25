"""Export evaluation summaries, original reports and saved time series to CSV/Excel."""
import argparse
import csv
import math
from pathlib import Path
import re

DEFAULT_DIRECTORIES = ['results', 'results_dongchen_baseline', 'results_zmerge', 'results_nash_double_q']
NUMBER = r'[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?|[-+]?(?:nan|inf)'
HEADER = re.compile(r'^(.*?):\s*(\d+(?:\.\d+)?)\s*(?:veh/h/lane|vph)\s*\|\s*'
                    r'(\d+(?:\.\d+)?)%\s*PR\s*\(Over\s*(\d+)\s*Seeds\)', re.M)
METRIC = re.compile(r'^\s*([^:\n]+):\s*('+NUMBER+r')(?:\s*±\s*('+NUMBER+r'))?\s*([^\n]*)$', re.M | re.I)
# Keep the established CSV column names for the existing plotting programs.
COLUMNS = {
    'Equivalent Hourly Throughput': 'Equivalent Hourly Throughput (veh/h)',
    'Raw Arrivals': 'Raw Arrivals (veh)', 'Mean Speed': 'Mean Speed (m/s)',
    'Mean Travel Time': 'Mean Travel Time (s)', 'Collision Rate': 'Collision Rate (%)',
    'Avg Near-Collisions': 'Avg Near-Collisions', 'Avg TTC Violations': 'Avg TTC Violations',
    'Avg Hard Braking': 'Avg Hard Braking', 'Acceleration Var': 'Acceleration Var (m/s^2)',
    'Mean Abs Jerk': 'Mean Abs Jerk (m/s^3)', 'Wave Intensity': 'Wave Intensity (Var)',
    'Raw Interventions': 'Raw Interventions (Sum)', 'Interventions Per CAV': 'Interventions Per CAV',
    'Duty Cycle': 'Duty Cycle (%)', 'Mean Decel Magnitude': 'Mean Decel Magnitude (m/s^2)',
    'Mean Duration': 'Mean Duration (s)', 'Success Rate': 'Success Rate (%)',
    'Avg Max Queue': 'Avg Max Queue (veh)', 'Absolute Max Queue': 'Absolute Max Queue (veh)',
    'Mean Delay': 'Mean Delay (s)', '90th-Percentile Delay': '90th-Percentile Delay (s)',
    'Traffic Efficiency': 'Mean Speed (m/s)', 'Traffic Safety': 'Collision Rate (%)',
    'Driving Comfort': 'Acceleration Var (m/s^2)',
    'Average Queue Length': 'Average Queue Length (veh)', 'Merging Success Rate': 'Success Rate (%)',
}


def discover_reports(directories):
    files = set()
    for directory in directories:
        path = Path(directory).expanduser()
        candidates = [path] if path.is_file() else path.rglob('*.txt')
        for file in candidates:
            if file.name.startswith('report_') or file.name.endswith('evaluation_results.txt'):
                files.add(file.resolve())
    return sorted(files)


def parse_report(path):
    text = path.read_text(encoding='utf-8-sig', errors='replace')
    header = HEADER.search(text)
    row = {'Model': 'Z-MERGE' if 'FINAL Z-MERGE' in text else 'UNRECOGNIZED',
           'Demand (vph/lane)': None, 'Penetration Rate (%)': None, 'Seeds Tested': None}
    if header:
        row.update({'Model': header[1].strip(), 'Demand (vph/lane)': float(header[2]),
                    'Penetration Rate (%)': float(header[3]), 'Seeds Tested': int(header[4])})
    metrics = []
    for match in METRIC.finditer(text):
        if header and header.start() <= match.start() < header.end():
            continue
        label, value, ci, unit = match.groups()
        label = label.strip()
        base = label.split(' (')[0]
        column = COLUMNS.get(base, label + (f' ({unit.strip()})' if unit.strip() else ''))
        number, interval = float(value), float(ci) if ci else None
        row[column] = f'{value} ± {ci}' if ci else number
        metrics.append({'Metric': column, 'Reported label': label, 'Value': number,
                        '95% CI': interval, 'Unit': unit.strip()})
    latency = re.search(r'remains below\s*('+NUMBER+r')\s*ms.*?average delay of\s*('+NUMBER+r')\s*ms', text, re.S)
    if latency:
        for column, value in zip(['Inference P90 (ms)', 'Inference Mean (ms)'], latency.groups()):
            row[column] = float(value)
            metrics.append({'Metric': column, 'Reported label': column, 'Value': float(value),
                            '95% CI': None, 'Unit': 'ms'})
    row['Source report'] = str(path)
    return row, metrics, text


def collect_reports(directories):
    reports = [(path, *parse_report(path)) for path in discover_reports(directories)]
    if not reports:
        raise ValueError('No evaluation reports found in the selected paths.')
    return reports


def write_csv(reports, output_csv):
    columns = list(dict.fromkeys(['Model', 'Demand (vph/lane)', 'Penetration Rate (%)', 'Seeds Tested']
                                + list(COLUMNS.values())
                                + [key for _, row, _, _ in reports for key in row]))
    output_csv = Path(output_csv)
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(row for _, row, _, _ in reports)
    return columns


def compile_reports_to_csv(search_directories, output_csv='combined_traffic_results.csv'):
    """Preserve the original callable API for existing analysis users."""
    paths = [search_directories] if isinstance(search_directories, (str, Path)) else search_directories
    reports = collect_reports(paths)
    write_csv(reports, output_csv)
    print(f'Compiled {len(reports)} reports into {output_csv}')
    return [row for _, row, _, _ in reports]


def export_results(directories, output_csv, output_xlsx, include_arrays=True):
    import numpy as np
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    reports = collect_reports(directories)
    columns = write_csv(reports, output_csv)
    book = Workbook(write_only=True)

    def append(sheet, values):
        cells = []
        for value in values:
            if isinstance(value, float) and not math.isfinite(value):
                value = str(value)
            cell = WriteOnlyCell(sheet, value=value)
            # Report text and filenames are data, including strings starting '='.
            if isinstance(value, str):
                cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)

    summary = book.create_sheet('summary')
    append(summary, columns)
    metrics = book.create_sheet('metrics')
    append(metrics, ['Model', 'Demand (vph/lane)', 'Penetration Rate (%)', 'Seeds Tested',
                     'Metric', 'Reported label', 'Value', '95% CI', 'Unit', 'Source report'])
    raw = book.create_sheet('reports')
    append(raw, ['Source report', 'Text part', 'Original report text'])
    for path, row, values, text in reports:
        append(summary, [row.get(key) for key in columns])
        for metric in values:
            append(metrics, [row[key] for key in ['Model', 'Demand (vph/lane)', 'Penetration Rate (%)', 'Seeds Tested']]
                   + [metric[key] for key in ['Metric', 'Reported label', 'Value', '95% CI', 'Unit']]+[str(path)])
        # Excel has a 32,767-character cell limit. Preserve long reports in parts.
        for part, start in enumerate(range(0, len(text), 32000), 1):
            append(raw, [str(path), part, text[start:start+32000]])

    artifacts = book.create_sheet('artifacts')
    append(artifacts, ['Path', 'Type', 'Bytes', 'Array shape', 'Array dtype'])
    array_sheet = book.create_sheet('arrays')
    array_header = ['Source array', 'Index (zero-based)', 'Value']
    append(array_sheet, array_header)
    array_rows, sheet_number = 1, 1
    files = set()
    for directory in directories:
        path = Path(directory)
        if path.is_dir():
            files.update(p.resolve() for p in path.rglob('*') if p.suffix.lower() in {'.png', '.mp4', '.npy'})
    for path in sorted(files):
        evaluation_array = path.suffix == '.npy' and path.name.startswith(('speed_array_', 'decel_array_', 'jerk_array_'))
        array = np.load(path, mmap_mode='r', allow_pickle=False) if evaluation_array else None
        append(artifacts, [str(path), path.suffix, path.stat().st_size,
                           str(array.shape) if array is not None else None,
                           str(array.dtype) if array is not None else None])
        if include_arrays and array is not None:
            for index, value in np.ndenumerate(array):
                if array_rows == 1048576:
                    sheet_number += 1
                    array_sheet = book.create_sheet(f'arrays_{sheet_number}')
                    append(array_sheet, array_header)
                    array_rows = 1
                append(array_sheet, [str(path), ','.join(map(str, index)), value.item()])
                array_rows += 1
    notes = book.create_sheet('notes')
    for line in [
        'summary: one row per evaluation report; missing/unreported values remain blank.',
        'metrics: numeric values and separate confidence intervals; reports: original text.',
        'arrays: all saved evaluation speed/deceleration/jerk values unless --no-arrays was selected.',
        'artifacts: paths to arrays, figures and recordings; binary images/videos are not embedded.',
        'Z-Merge reports omit demand, PR, seeds and confidence intervals; these are not inferred.',
        'Average queue length (Z-Merge) differs from average maximum queue length (other methods).',
        'Reports from different source directories are retained as separate runs, not averaged.',
        'No missing raw per-seed data can be reconstructed from an aggregate report.',
    ]:
        append(notes, [line])
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_xlsx)
    print(f'Exported {len(reports)} reports to {output_csv} and {output_xlsx}')
    return len(reports)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--search-dirs', nargs='+', default=DEFAULT_DIRECTORIES)
    parser.add_argument('--output-csv', default='combined_traffic_results.csv')
    parser.add_argument('--output-xlsx', default='combined_traffic_results.xlsx')
    parser.add_argument('--no-arrays', action='store_true', help='Index array files without embedding their values')
    args = parser.parse_args()
    try:
        export_results(args.search_dirs, args.output_csv, args.output_xlsx, not args.no_arrays)
    except ValueError as error:
        parser.error(str(error))


if __name__ == '__main__':
    main()
