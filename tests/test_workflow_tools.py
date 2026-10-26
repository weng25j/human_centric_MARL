"""Check spreadsheet contents and orchestration without long experiment runs."""
import csv
import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from analysis.combine_evaluation_reports import collect_reports, export_results

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('pipeline', ROOT/'scripts/run_experiments.py')
pipeline = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pipeline)


class SpreadsheetTests(unittest.TestCase):
    def test_both_report_formats_and_arrays_are_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root/'report_shepherd.txt').write_text('TRAFFIC SHEPHERD: 800 veh/h/lane | 10% PR (Over 20 Seeds)\n'
                'Mean Speed: 12.50 ± 0.75 m/s\nAvg Max Queue: 7.00 ± 1.00 vehicles\n', encoding='utf-8')
            (root/'run_evaluation_results.txt').write_text('FINAL Z-MERGE METHODOLOGY EVALUATION RESULTS\n'
                'Traffic Efficiency (Avg Speed): 15.01 m/s\nAverage Queue Length: 1.78 vehicles/step\n', encoding='utf-8')
            np.save(root/'speed_array_SHEPHERD_800vph_10PR.npy', np.array([1.25, 2.5]))
            count = export_results([root, root], root/'out.csv', root/'out.xlsx')
            self.assertEqual(count, 2)
            book = load_workbook(root/'out.xlsx', read_only=True, data_only=True)
            summary = list(book['summary'].values)
            rows = [dict(zip(summary[0], row)) for row in summary[1:]]
            shepherd = next(row for row in rows if row['Model'] == 'TRAFFIC SHEPHERD')
            zmerge = next(row for row in rows if row['Model'] == 'Z-MERGE')
            self.assertEqual(shepherd['Mean Speed (m/s)'], '12.50 ± 0.75')
            self.assertIsNone(zmerge['Demand (vph/lane)'])
            self.assertEqual(zmerge['Average Queue Length (veh)'], 1.78)
            self.assertIsNone(zmerge['Avg Max Queue (veh)'])
            self.assertEqual([r[2] for r in list(book['arrays'].values)[1:]], [1.25, 2.5])
            metrics = list(book['metrics'].values)
            self.assertEqual(len(metrics), 5)  # header and four metrics, no duplicate report header
            with (root/'out.csv').open(encoding='utf-8-sig') as stream:
                self.assertEqual(len(list(csv.DictReader(stream))), 2)
            book.close()

    def test_missing_reports_fail_instead_of_creating_empty_success(self):
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaisesRegex(ValueError, 'No evaluation reports'):
                collect_reports([td])

    def test_report_text_is_not_an_excel_formula(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root/'report_example.txt').write_text('=FAKE: 800 veh/h/lane | 10% PR (Over 1 Seeds)\nMean Speed: 1 ± 0 m/s', encoding='utf-8')
            export_results([root], root/'out.csv', root/'out.xlsx')
            book = load_workbook(root/'out.xlsx')
            self.assertEqual(book['summary']['A2'].data_type, 's')
            book.close()


class PipelineTests(unittest.TestCase):
    def test_video_options_reach_common_and_v2_trainers_and_evaluators(self):
        for selection in (['--suite', 'common'], ['--methods', 'shepherd_v2']):
            with tempfile.TemporaryDirectory() as td, patch('builtins.print') as output:
                pipeline.main([*selection, '--dry-run', '--run-dir', td, '--record-training',
                               '--record-evaluation', '--video-runs', '2', '--video-max-steps', '120', '--video-every', '2'])
                commands = [str(c.args[0]) for c in output.call_args_list if ' -m training.' in str(c) or ' -m evaluation.' in str(c)]
                self.assertTrue(commands)
                self.assertTrue(all('--record-video --video-runs 2 --video-max-steps 120 --video-every 2' in c for c in commands))
        with patch('sys.stderr'), self.assertRaises(SystemExit):
            pipeline.parse_args(['--record-training'])

    def test_common_suite_includes_every_method_and_keeps_legacy_rewards_separate(self):
        args = pipeline.parse_args(['--suite', 'common'])
        self.assertEqual(set(args.methods), set(pipeline.COMMON_METHODS))
        with tempfile.TemporaryDirectory() as td, patch('builtins.print') as output:
            destination = Path(td)/'not_created'
            pipeline.main(['--suite', 'common', '--dry-run', '--run-dir', str(destination)])
            text = '\n'.join(str(c) for c in output.call_args_list)
            for method in pipeline.COMMON_METHODS:
                self.assertIn(method, text)
            self.assertIn('analysis.compare_common_reward', text)
            self.assertNotIn('training.train_traffic_shepherd ', text)
            self.assertFalse(destination.exists())
        with patch('sys.stderr'), self.assertRaises(SystemExit):
            pipeline.parse_args(['--suite', 'common', '--methods', 'shepherd'])

    def test_v2_is_opt_in_and_has_separate_modules_and_model_location(self):
        self.assertNotIn('shepherd_v2', pipeline.parse_args([]).methods)
        with tempfile.TemporaryDirectory() as td, patch('builtins.print') as output:
            destination = Path(td)/'not_created'
            args = pipeline.parse_args(['--dry-run', '--methods', 'shepherd_v2', '--run-dir', str(destination)])
            self.assertEqual(pipeline.model_location(args, 'shepherd_v2', destination),
                             destination/'models/shepherd_v2/models')
            pipeline.main(['--dry-run', '--methods', 'shepherd_v2', '--run-dir', str(destination),
                           '--v2-episodes', '1', '--v2-train-steps', '500', '--v2-eval-seeds', '0,25'])
            text = '\n'.join(str(c) for c in output.call_args_list)
            self.assertIn('training.train_traffic_shepherd_v2', text)
            self.assertIn('evaluation.evaluate_traffic_shepherd_v2', text)
            self.assertIn('analysis.combine_evaluation_reports', text)
            self.assertNotIn('training.train_traffic_shepherd ', text)
            self.assertFalse(destination.exists())

    def test_dry_run_lists_all_methods_and_export_without_writing(self):
        with tempfile.TemporaryDirectory() as td, patch('builtins.print') as output:
            destination = Path(td)/'not_created'
            pipeline.main(['--dry-run', '--run-dir', str(destination)])
            text = '\n'.join(str(c) for c in output.call_args_list)
            for module in ['training.train_traffic_shepherd', 'training.train_flat_marl', 'training.train_z_merge',
                           'evaluation.evaluate_traffic_shepherd', 'evaluation.evaluate_flat_marl',
                           'evaluation.evaluate_z_merge', 'evaluation.evaluate_rule_based',
                           'training.train_nash_double_q', 'evaluation.evaluate_nash_double_q',
                           'analysis.combine_evaluation_reports']:
                self.assertIn(module, text)
            self.assertFalse(destination.exists())

    def test_invalid_mapppo_checkpoint_cannot_be_evaluated(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root/'checkpoint-not-a-number.pt').touch()
            with self.assertRaises(FileNotFoundError):
                pipeline.validate_weights('shepherd', root)
            (root/'checkpoint-50.pt').touch()
            pipeline.validate_weights('shepherd', root)


if __name__ == '__main__':
    unittest.main()
